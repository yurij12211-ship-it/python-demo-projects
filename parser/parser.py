"""
Парсер каталогу books.toscrape.com (навчальний сайт для скрапінгу).
Збирає назву, ціну, рейтинг, наявність і посилання з усіх сторінок
і зберігає результат у books.xlsx з фільтрами та закріпленою шапкою.

Запуск:  pip install requests beautifulsoup4 openpyxl
         python parser.py
"""
import time
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

BASE = "https://books.toscrape.com/catalogue/"
RATING = {"One": 1, "Two": 2, "Three": 3, "Four": 4, "Five": 5}
HEADERS = {"User-Agent": "Mozilla/5.0 (demo-parser)"}


def parse_page(url: str) -> tuple[list[dict], str | None]:
    resp = requests.get(url, headers=HEADERS, timeout=15)
    resp.raise_for_status()
    soup = BeautifulSoup(resp.text, "html.parser")
    rows = []
    for card in soup.select("article.product_pod"):
        link = card.select_one("h3 a")
        rows.append({
            "title": link["title"],
            "price": float(card.select_one(".price_color").text.strip().lstrip("£Â")),
            "rating": RATING.get(card.select_one(".star-rating")["class"][1], 0),
            "in_stock": "In stock" in card.select_one(".availability").text,
            "url": urljoin(url, link["href"]),
        })
    nxt = soup.select_one("li.next a")
    return rows, urljoin(url, nxt["href"]) if nxt else None


def save_xlsx(rows: list[dict], path: str = "books.xlsx") -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Книги"
    head = ["Назва", "Ціна, £", "Рейтинг", "В наявності", "Посилання"]
    ws.append(head)
    for r in rows:
        ws.append([r["title"], r["price"], r["rating"],
                   "так" if r["in_stock"] else "ні", r["url"]])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F6F50")
        c.alignment = Alignment(horizontal="center")
    for col, width in zip("ABCDE", (55, 10, 10, 13, 70)):
        ws.column_dimensions[col].width = width
    for cell in ws["B"][1:]:
        cell.number_format = "0.00"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def main(max_pages: int | None = None) -> None:
    url, rows, page = BASE + "page-1.html", [], 0
    while url and (max_pages is None or page < max_pages):
        part, url = parse_page(url)
        rows += part
        page += 1
        print(f"сторінка {page}: +{len(part)} (усього {len(rows)})")
        time.sleep(0.5)  # ввічлива пауза між запитами
    save_xlsx(rows)
    print(f"Готово: {len(rows)} рядків у books.xlsx")


if __name__ == "__main__":
    main()
